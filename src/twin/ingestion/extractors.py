"""Format-specific text extraction from source files."""

import csv
import email
import json
from dataclasses import dataclass, field
from email import policy
from pathlib import Path

import docx
import pdfplumber


@dataclass
class ExtractionResult:
    text: str
    metadata: dict = field(default_factory=dict)
    file_type: str = ""


def detect_file_type(path: Path) -> str:
    suffix = path.suffix.lower()
    type_map = {
        ".pdf": "pdf",
        ".docx": "docx",
        ".doc": "doc",
        ".eml": "email",
        ".msg": "msg",
        ".md": "md",
        ".markdown": "md",
        ".txt": "txt",
        ".json": "json",
        ".xlsx": "xlsx",
        ".xls": "xlsx",
        ".pptx": "pptx",
        ".ppt": "ppt",
        ".csv": "csv",
        ".png": "image",
        ".jpg": "image",
        ".jpeg": "image",
    }
    return type_map.get(suffix, "txt")


def extract(path: Path) -> ExtractionResult:
    """Extract text from a file, dispatching by format."""
    file_type = detect_file_type(path)
    extractors = {
        "pdf": _extract_pdf,
        "docx": _extract_docx,
        "doc": _extract_doc,
        "email": _extract_email,
        "msg": _extract_msg,
        "md": _extract_plaintext,
        "txt": _extract_plaintext,
        "json": _extract_json,
        "xlsx": _extract_xlsx,
        "pptx": _extract_pptx,
        "ppt": _extract_ppt,
        "csv": _extract_csv,
        "image": _extract_image,
    }
    extractor = extractors.get(file_type, _extract_plaintext)
    result = extractor(path)
    result.file_type = file_type
    return result


def _extract_pdf(path: Path) -> ExtractionResult:
    pages = []
    metadata = {}
    with pdfplumber.open(path) as pdf:
        metadata["page_count"] = len(pdf.pages)
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
    return ExtractionResult(text="\n\n".join(pages), metadata=metadata)


def _extract_docx(path: Path) -> ExtractionResult:
    doc = docx.Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    metadata = {}
    core = doc.core_properties
    if core.author:
        metadata["author"] = core.author
    if core.title:
        metadata["title"] = core.title
    if core.created:
        metadata["created"] = core.created.isoformat()
    return ExtractionResult(text="\n\n".join(paragraphs), metadata=metadata)


def _extract_email(path: Path) -> ExtractionResult:
    raw = path.read_bytes()
    msg = email.message_from_bytes(raw, policy=policy.default)

    metadata = {
        "subject": msg["subject"] or "",
        "from": msg["from"] or "",
        "to": msg["to"] or "",
        "date": msg["date"] or "",
        "cc": msg["cc"] or "",
        "bcc": msg["bcc"] or "",
        "message_id": msg["message-id"] or "",
        "in_reply_to": msg["in-reply-to"] or "",
        "references": msg["references"] or "",
        "thread_topic": msg["thread-topic"] or "",
        "thread_index": msg["thread-index"] or "",
    }

    body = msg.get_body(preferencelist=("plain", "html"))
    if body:
        content_type = body.get_content_type()
        body_text = body.get_content()
        if content_type == "text/html":
            body_text = _html_to_text(body_text)
    else:
        body_text = ""

    thread_messages = _parse_thread(body_text)
    metadata["thread_depth"] = len(thread_messages)
    metadata["is_reply"] = bool(msg["in-reply-to"])

    attachments = _extract_attachments(msg, path)
    if attachments:
        metadata["attachments"] = attachments
        metadata["attachment_count"] = len(attachments)

    sections = []
    sections.append(f"Subject: {metadata['subject']}")
    sections.append(f"From: {metadata['from']}")
    sections.append(f"To: {metadata['to']}")
    if metadata["cc"]:
        sections.append(f"CC: {metadata['cc']}")
    sections.append(f"Date: {metadata['date']}")
    sections.append("")

    if len(thread_messages) > 1:
        for i, tm in enumerate(thread_messages):
            label = "Latest Reply" if i == 0 else f"Previous Message ({i})"
            sections.append(f"--- {label} ---")
            if tm.get("from"):
                sections.append(f"From: {tm['from']}")
            if tm.get("date"):
                sections.append(f"Date: {tm['date']}")
            sections.append(tm["body"])
            sections.append("")
    else:
        sections.append(body_text.strip())

    if attachments:
        sections.append("")
        sections.append("--- Attachments ---")
        for att in attachments:
            sections.append(
                f"  [{att['filename']}] ({att['content_type']}, {att['size_bytes']} bytes)"
            )
            if att.get("saved_path"):
                sections.append(f"    Saved: {att['saved_path']}")

    text = "\n".join(sections)
    return ExtractionResult(text=text.strip(), metadata=metadata)


def _html_to_text(html: str) -> str:
    """Convert HTML email body to plain text."""
    import re

    text = re.sub(r"<br\s*/?>", "\n", html, flags=re.IGNORECASE)
    text = re.sub(r"<p[^>]*>", "\n", text, flags=re.IGNORECASE)
    text = re.sub(r"</p>", "\n", text, flags=re.IGNORECASE)
    text = re.sub(r"<div[^>]*>", "\n", text, flags=re.IGNORECASE)
    text = re.sub(r"<[^>]+>", "", text)
    import html as html_module

    text = html_module.unescape(text)
    lines = [line.strip() for line in text.splitlines()]
    text = "\n".join(lines)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _parse_thread(body_text: str) -> list[dict]:
    """Split an email body into individual messages in a thread.

    Returns a list of dicts with keys: from, date, body.
    First element is the most recent message.
    """
    import re

    separator_pattern = re.compile(
        r"^(?:"
        r"[-_]{2,}\s*(?:Original Message|Forwarded message)?\s*[-_]*"
        r"|From:\s+.+?(?:\n|$)"
        r"|On .+? wrote:"
        r")",
        re.MULTILINE | re.IGNORECASE,
    )

    outlook_from_pattern = re.compile(
        r"^From:\s*(.+?)$\s*^Sent:\s*(.+?)$\s*^To:",
        re.MULTILINE,
    )

    messages = []
    parts = separator_pattern.split(body_text)

    if len(parts) <= 1:
        return [{"from": "", "date": "", "body": body_text.strip()}]

    for i, part in enumerate(parts):
        part = part.strip()
        if not part:
            continue

        msg_from = ""
        msg_date = ""

        from_match = outlook_from_pattern.search(part)
        if from_match:
            msg_from = from_match.group(1).strip()
            msg_date = from_match.group(2).strip()
            part = part[from_match.end() :].strip()
            to_line = re.match(r"^To:.*$", part, re.MULTILINE)
            if to_line:
                part = part[to_line.end() :].strip()
            subject_line = re.match(r"^Subject:.*$", part, re.MULTILINE)
            if subject_line:
                part = part[subject_line.end() :].strip()

        if part:
            messages.append({"from": msg_from, "date": msg_date, "body": part})

    return messages if messages else [{"from": "", "date": "", "body": body_text.strip()}]


def _extract_attachments(msg: email.message.Message, eml_path: Path) -> list[dict]:
    """Extract attachments and inline images from an email message.

    Saves non-trivial attachments to data/email_attachments/<email-stem>/
    """
    attachments: list[dict] = []
    attachments_dir = Path("data/email_attachments") / eml_path.stem

    for part in msg.walk():
        disposition = part.get_content_disposition()
        content_type = part.get_content_type()
        maintype = part.get_content_maintype()

        is_attachment = disposition == "attachment"
        is_inline_media = disposition == "inline" and maintype in ("image", "application")

        if not (is_attachment or is_inline_media):
            continue

        filename = (
            part.get_filename() or f"unnamed_{len(attachments)}.{_guess_extension(content_type)}"
        )
        payload = part.get_payload(decode=True)
        if not payload:
            continue

        att_info = {
            "filename": filename,
            "content_type": content_type,
            "size_bytes": len(payload),
            "disposition": disposition or "inline",
            "content_id": part.get("Content-ID", ""),
        }

        if len(payload) > 1024:
            attachments_dir.mkdir(parents=True, exist_ok=True)
            save_path = attachments_dir / filename
            save_path.write_bytes(payload)  # type: ignore[arg-type]
            att_info["saved_path"] = str(save_path)

        attachments.append(att_info)

    return attachments


def _guess_extension(content_type: str) -> str:
    ext_map = {
        "image/png": "png",
        "image/jpeg": "jpg",
        "image/gif": "gif",
        "application/pdf": "pdf",
        "application/octet-stream": "bin",
    }
    return ext_map.get(content_type, "bin")


def _extract_msg(path: Path) -> ExtractionResult:
    """Extract from Outlook .msg binary format using extract-msg."""
    import extract_msg

    msg = extract_msg.Message(str(path))

    metadata = {
        "subject": msg.subject or "",
        "from": msg.sender or "",
        "to": msg.to or "",
        "date": msg.date.isoformat() if msg.date else "",
        "cc": msg.cc or "",
        "bcc": msg.bcc or "",
        "message_id": msg.messageId or "",
        "in_reply_to": msg.inReplyTo or "",
        "thread_topic": msg.subject or "",
        "is_reply": bool(msg.inReplyTo),
    }

    body_text = msg.body or ""
    if not body_text and msg.htmlBody:
        body_text = _html_to_text(
            msg.htmlBody.decode("utf-8", errors="replace")
            if isinstance(msg.htmlBody, bytes)
            else msg.htmlBody
        )

    thread_messages = _parse_thread(body_text)
    metadata["thread_depth"] = len(thread_messages)

    attachments: list[dict] = []
    attachments_dir = Path("data/email_attachments") / path.stem
    for att in msg.attachments:  # type: ignore[union-attr]
        data = getattr(att, "data", None)
        att_info: dict = {
            "filename": getattr(att, "longFilename", None)
            or getattr(att, "shortFilename", None)
            or f"unnamed_{len(attachments)}.bin",
            "content_type": getattr(att, "mimetype", None) or "application/octet-stream",
            "size_bytes": len(data) if data else 0,
            "disposition": "attachment",
        }
        if data and len(data) > 1024:
            attachments_dir.mkdir(parents=True, exist_ok=True)
            save_path = attachments_dir / att_info["filename"]
            save_path.write_bytes(data)
            att_info["saved_path"] = str(save_path)
        attachments.append(att_info)

    if attachments:
        metadata["attachments"] = attachments
        metadata["attachment_count"] = len(attachments)

    sections = []
    sections.append(f"Subject: {metadata['subject']}")
    sections.append(f"From: {metadata['from']}")
    sections.append(f"To: {metadata['to']}")
    if metadata["cc"]:
        sections.append(f"CC: {metadata['cc']}")
    sections.append(f"Date: {metadata['date']}")
    sections.append("")

    if len(thread_messages) > 1:
        for i, tm in enumerate(thread_messages):
            label = "Latest Reply" if i == 0 else f"Previous Message ({i})"
            sections.append(f"--- {label} ---")
            if tm.get("from"):
                sections.append(f"From: {tm['from']}")
            if tm.get("date"):
                sections.append(f"Date: {tm['date']}")
            sections.append(tm["body"])
            sections.append("")
    else:
        sections.append(body_text.strip())

    if attachments:
        sections.append("")
        sections.append("--- Attachments ---")
        for att_item in attachments:  # type: ignore[assignment]
            name = att_item["filename"]
            ctype = att_item["content_type"]
            size = att_item["size_bytes"]
            sections.append(f"  [{name}] ({ctype}, {size} bytes)")
            if att_item.get("saved_path"):
                sections.append(f"    Saved: {att_item['saved_path']}")

    msg.close()
    text = "\n".join(sections)
    return ExtractionResult(text=text.strip(), metadata=metadata)


def _extract_plaintext(path: Path) -> ExtractionResult:
    try:
        text = path.read_text(encoding="utf-8").strip()
    except UnicodeDecodeError:
        text = path.read_text(encoding="latin-1").strip()
    return ExtractionResult(text=text)


def _extract_json(path: Path) -> ExtractionResult:
    raw = path.read_text(encoding="utf-8")
    data = json.loads(raw)
    metadata = {}
    if isinstance(data, dict):
        metadata = {k: v for k, v in data.items() if isinstance(v, str | int | float | bool)}
    return ExtractionResult(text=raw, metadata=metadata)


def _extract_xlsx(path: Path) -> ExtractionResult:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets_text = []
    metadata = {"sheet_count": len(wb.sheetnames), "sheets": wb.sheetnames}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            if any(cells):
                rows.append("\t".join(cells))
        if rows:
            sheets_text.append(f"[Sheet: {sheet_name}]\n" + "\n".join(rows))
    wb.close()
    return ExtractionResult(text="\n\n".join(sheets_text), metadata=metadata)


def _extract_pptx(path: Path) -> ExtractionResult:
    from pptx import Presentation

    prs = Presentation(str(path))
    slides_text = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        texts.append(text)
        if texts:
            slides_text.append(f"[Slide {i}]\n" + "\n".join(texts))
    metadata = {"slide_count": len(prs.slides)}
    return ExtractionResult(text="\n\n".join(slides_text), metadata=metadata)


def _extract_ppt(path: Path) -> ExtractionResult:
    """Extract text from legacy .ppt files using filename context as fallback."""
    # python-pptx doesn't support .ppt (legacy binary format)
    # Use the filename as the extracted content for classification
    return ExtractionResult(
        text=f"[Legacy PPT file: {path.name}]",
        metadata={"format": "legacy_ppt", "note": "Binary format, text from filename only"},
    )


def _extract_doc(path: Path) -> ExtractionResult:
    """Extract text from legacy .doc files using filename context as fallback."""
    # python-docx doesn't support .doc (legacy binary format)
    return ExtractionResult(
        text=f"[Legacy DOC file: {path.name}]",
        metadata={"format": "legacy_doc", "note": "Binary format, text from filename only"},
    )


def _extract_csv(path: Path) -> ExtractionResult:
    rows_text = []
    with open(path, newline="", encoding="utf-8", errors="replace") as f:
        reader = csv.reader(f)
        for row in reader:
            if any(cell.strip() for cell in row):
                rows_text.append("\t".join(row))
    return ExtractionResult(text="\n".join(rows_text), metadata={"row_count": len(rows_text)})


def _extract_image(path: Path) -> ExtractionResult:
    """Extract text from images using Tesseract OCR."""
    try:
        import pytesseract
        from PIL import Image

        img = Image.open(path)
        text = pytesseract.image_to_string(img).strip()
        if text:
            return ExtractionResult(
                text=text, metadata={"ocr": True, "dimensions": f"{img.width}x{img.height}"}
            )
    except Exception:
        pass
    return ExtractionResult(
        text=f"[Image file: {path.name}]",
        metadata={"ocr": False, "note": "OCR failed or no text detected"},
    )
