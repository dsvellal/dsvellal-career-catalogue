#!/usr/bin/env python3
"""
Process 92 Philips session feedback Excel files into structured JSON.
Extracts ratings, text feedback, dates, and categorizes each session.
"""

import argparse
import json
import os
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from statistics import mean

import openpyxl

FEEDBACK_DIR = Path("/Users/dsvellal/Downloads/Website/Philips-Sessions-Feedback")
OUTPUT_DIR = Path(__file__).resolve().parent

# Month name to number mapping
MONTH_MAP = {
    'jan': 1, 'january': 1, 'feb': 2, 'february': 2, 'mar': 3, 'march': 3,
    'apr': 4, 'april': 4, 'may': 5, 'jun': 6, 'june': 6,
    'jul': 7, 'july': 7, 'aug': 8, 'august': 8, 'sep': 9, 'sept': 9, 'september': 9,
    'oct': 10, 'october': 10, 'nov': 11, 'november': 11, 'dec': 12, 'december': 12
}


def extract_date_from_filename(filename):
    """Extract date from filename using various patterns."""
    name = filename.replace('.xlsx', '')

    # Pattern: YYYYMMDD at start like 20190415
    m = re.search(r'(\d{4})(\d{2})(\d{2})', name)
    if m and int(m.group(1)) >= 2018 and int(m.group(1)) <= 2026:
        year, month, day = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if 1 <= month <= 12 and 1 <= day <= 31:
            return f"{year}-{month:02d}-{day:02d}"

    # Pattern: [Feedback] 20250312 ...
    m = re.match(r'\[Feedback\]\s*(\d{4})(\d{2})(\d{2})', name)
    if m:
        year, month, day = int(m.group(1)), int(m.group(2)), int(m.group(3))
        return f"{year}-{month:02d}-{day:02d}"

    # Pattern: DDth Month YYYY or DDth Month Year
    month_pattern = (
        r"(\d{1,2})(?:st|nd|rd|th)?\s+"
        r"(Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|"
        r"Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|Nov(?:ember)?|"
        r"Dec(?:ember)?)\s+(\d{4})"
    )
    m = re.search(month_pattern, name, re.IGNORECASE)
    if m:
        day = int(m.group(1))
        month = MONTH_MAP.get(m.group(2).lower()[:3], 0)
        year = int(m.group(3))
        if month and 1 <= day <= 31:
            return f"{year}-{month:02d}-{day:02d}"

    # Pattern: Dec-16-Workshop, Dec-17-Workshop, Dec-18
    m = re.search(
        r"(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)-(\d{1,2})",
        name,
        re.IGNORECASE,
    )
    if m:
        month = MONTH_MAP.get(m.group(1).lower()[:3], 0)
        day = int(m.group(2))
        # Guess year from context - these Dec workshops are from 2024
        if 'Workshop' in name or 'Feedback' in name:
            return f"2024-{month:02d}-{day:02d}"

    # Pattern: just a year like "2018 -" or "2019 " or "2020 -"
    m = re.match(r'(\d{4})\s*[-.]', name)
    if m:
        year = int(m.group(1))
        if 2018 <= year <= 2026:
            return str(year)

    # Pattern: year somewhere in the name for older files
    m = re.search(r'20(18|19|20|21|22|23|24|25|26)', name)
    if m:
        year = int("20" + m.group(1))
        # Only return year if we can't get more specific
        return str(year)

    return "unknown"


def categorize_session(filename, questions=None):
    """Categorize session based on filename and content."""
    name = filename.lower()

    # Before-after comparison
    if 'jscpd before' in name or 'jscpd after' in name:
        return "before-after"

    # Connects / personal feedback about Datta
    if '.connects' in name or 'connects' in name:
        return "connects"
    if (
        "feedback about datta" in name
        or "feedback for datta" in name
        or "feedback to datta" in name
    ):
        return "feedback-about-datta"
    if 'new hire mentoring' in name:
        return "feedback-about-datta"
    if 'pair.?programming.*feedback' in name or 'pair programming' in name:
        return "feedback-about-datta"
    if "feedback to dattaterya" in name or "feedback about the swcoe" in name:
        return "feedback-about-datta"
    if "tune observation" in name:
        return "feedback-about-datta"

    # Interview sessions
    if 'interview' in name or 'hiring' in name or 'panelist' in name:
        return "interview"

    # CoPilot / AI / GenAI
    if 'copilot' in name or 'co-pilot' in name or 'ai ' in name or 'ai-' in name or 'genai' in name:
        return "ai-genai"
    if 'context engineering' in name or 'agents' in name or 'elevate' in name:
        return "ai-genai"
    if 'prompt engineering' in name:
        return "ai-genai"

    # DORA
    if 'dora' in name or 'continuous value delivery' in name:
        return "dora"

    # CodeScene / Tech Debt
    if 'codescene' in name or 'tech debt' in name or 'tech-debt' in name:
        return "tech-debt"

    # Code quality sessions
    if 'code quality' in name or 'clean.?code' in name or 'code dojo' in name:
        return "code-quality"
    if 'code duplication' in name or 'jscpd' in name:
        return "code-quality"
    if 'unit test' in name or 'mutation test' in name or 'bdd' in name:
        return "code-quality"
    if 'behavior-driven' in name or 'best practices' in name:
        return "code-quality"
    if 'code review' in name or 'shift left' in name or 'back to basics' in name:
        return "code-quality"
    if 'i will code' in name or 'crafting code' in name:
        return "code-quality"

    # DORA / metrics
    if 'program behavior metrics' in name or 'observability' in name:
        return "dora"

    # Surveys / audience understanding (pre-session)
    if 'understanding' in name and 'audience' in name:
        return "other"
    if 'survey' in name and 'requirement' in name:
        return "other"

    # Other specific categories
    if 'microservices' in name or 'reliability' in name:
        return "code-quality"
    if 'traceability' in name:
        return "code-quality"
    if 'bar_skill raiser' in name or 'bar raiser' in name:
        return "interview"
    if 'github foundational' in name:
        return "ai-genai"
    if '62304' in name:
        return "code-quality"
    if 'analytics community' in name:
        return "ai-genai"
    if 'gdp engagement' in name or 'software excellence' in name or 'business interaction' in name:
        return "other"
    if 'grow' in name:
        return "ai-genai"

    return "other"


def is_numeric_rating(value, max_scale=10):
    """Check if a value is a numeric rating."""
    if value is None:
        return False
    try:
        v = float(value)
        return 1 <= v <= max_scale
    except (ValueError, TypeError):
        return False


def detect_rating_scale(values):
    """Detect whether ratings are on 1-5 or 1-10 scale."""
    numeric_vals = []
    for v in values:
        try:
            n = float(v)
            if 1 <= n <= 10:
                numeric_vals.append(n)
        except (ValueError, TypeError):
            continue

    if not numeric_vals:
        return None

    max_val = max(numeric_vals)
    if max_val <= 5:
        return 5
    return 10


def is_rating_column(header, values):
    """Determine if a column contains numeric ratings."""
    if header is None:
        return False

    header_str = str(header).lower()
    # Skip timestamp/ID columns
    skip_keywords = ['timestamp', 'id', 'email', 'name', 'date', 'time']
    for kw in skip_keywords:
        if kw in header_str:
            return False

    # Count numeric values in range
    numeric_count = 0
    total_non_empty = 0
    for v in values:
        if v is not None and str(v).strip():
            total_non_empty += 1
            try:
                n = float(v)
                if 1 <= n <= 10:
                    numeric_count += 1
            except (ValueError, TypeError):
                pass

    # If majority are numeric ratings
    if total_non_empty > 0 and numeric_count / total_non_empty >= 0.5:
        return True
    return False


def is_text_column(header, values):
    """Determine if a column contains meaningful text feedback."""
    if header is None:
        return False

    header_str = str(header).lower()
    # Skip timestamp/ID/metadata columns
    skip_keywords = ['timestamp', 'id', 'email address', 'start time', 'completion time',
                     'last modified', 'email id', 'repository link', 'what is your philips email',
                     'what is your code repository', 'command run']
    for kw in skip_keywords:
        if kw in header_str:
            return False

    # Check if values have meaningful text (not just numbers or short answers)
    text_count = 0
    for v in values:
        if v is not None:
            s = str(v).strip()
            if len(s) > 10 and not is_numeric_rating(v):
                text_count += 1

    return text_count > 0


def is_meaningful_feedback(text):
    """Filter out email content, metadata, and non-feedback text."""
    if not text or len(text) < 15:
        return False
    lower = text.lower()
    # Skip email-like content
    if '@philips.com' in lower and ('adding' in lower or 'cc' in lower or 'from:' in lower):
        return False
    # Skip URLs only
    if text.startswith('http') and ' ' not in text:
        return False
    # Skip entries that are just names/emails
    if '@' in text and len(text.split()) <= 5:
        return False
    return True


def extract_session_name(filename):
    """Extract a clean session name from the filename."""
    name = filename.replace('.xlsx', '')

    # Remove response count notation like (1-65)
    name = re.sub(r'\(\d+-\d+\)', '', name)

    # Remove [Feedback] prefix
    name = re.sub(r'\[.*?\]\s*', '', name)

    # Remove date prefixes like "DDth Month YYYY - "
    name = re.sub(r'\d{1,2}(?:st|nd|rd|th)?\s+\w+\s+\d{4}\s*-\s*', '', name)

    # Remove YYYYMMDD prefix
    name = re.sub(r'^\d{8}\s*-?\s*', '', name)

    # Remove " (1)" duplicate suffix
    name = re.sub(r'\s*\(\d+\)\s*$', '', name)

    # Remove trailing " - Feedback" or "- Session Feedback"
    name = re.sub(r'\s*-?\s*(?:Session\s+)?Feedback\s*$', '', name, flags=re.IGNORECASE)

    # Remove leading/trailing whitespace and dashes
    name = name.strip(' -_')

    return name if name else filename.replace('.xlsx', '')


def process_file(filepath):
    """Process a single Excel file and extract all data."""
    filename = os.path.basename(filepath)
    result = {
        "filename": filename,
        "session_name": extract_session_name(filename),
        "date": extract_date_from_filename(filename),
        "response_count": 0,
        "questions": [],
        "ratings": {},
        "text_feedback": [],
        "category": "",
        "is_pre_session_survey": False,
        "is_duplicate": False,
        "raw_data_sample": []
    }

    # Mark duplicates
    if filename.endswith(' (1).xlsx'):
        result["is_duplicate"] = True

    # Mark pre-session surveys
    if 'understanding' in filename.lower() and 'audience' in filename.lower():
        result["is_pre_session_survey"] = True

    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb.active

        # Read all rows using max_row/max_column for reliable reading
        rows = []
        for row in ws.iter_rows(
            min_row=1,
            max_row=ws.max_row,
            max_col=ws.max_column,
            values_only=True,
        ):
            rows.append(row)
        wb.close()

        if not rows:
            return result

        # First row is headers
        headers = list(rows[0])
        data_rows = rows[1:]

        result["questions"] = [str(h) if h is not None else "" for h in headers]
        result["response_count"] = len(data_rows)

        # Extract date from data if not found in filename
        if result["date"] == "unknown" and data_rows:
            # Look for timestamp columns
            for col_idx, header in enumerate(headers):
                timestamp_headers = ["start time", "completion time", "timestamp"]
                if header and any(
                    keyword in str(header).lower() for keyword in timestamp_headers
                ):
                    for row in data_rows:
                        if col_idx < len(row) and row[col_idx] is not None:
                            val = row[col_idx]
                            if hasattr(val, 'strftime'):
                                # It's a datetime object
                                result["date"] = val.strftime("%Y-%m-%d")
                                break
                            elif isinstance(val, str):
                                # Try parsing string date
                                timestamp_formats = [
                                    "%Y-%m-%d %H:%M:%S",
                                    "%m/%d/%Y %H:%M:%S",
                                    "%d/%m/%Y %H:%M:%S",
                                ]
                                for fmt in timestamp_formats:
                                    try:
                                        dt = datetime.strptime(val.split('.')[0], fmt)
                                        result["date"] = dt.strftime("%Y-%m-%d")
                                        break
                                    except ValueError:
                                        pass
                    if result["date"] != "unknown":
                        break

        # Process each column
        for col_idx, header in enumerate(headers):
            if header is None:
                continue

            # Get all values for this column
            col_values = [row[col_idx] if col_idx < len(row) else None for row in data_rows]

            # Check if it's a rating column
            if is_rating_column(header, col_values):
                numeric_vals = []
                for v in col_values:
                    try:
                        n = float(v)
                        if 1 <= n <= 10:
                            numeric_vals.append(n)
                    except (ValueError, TypeError):
                        pass

                if numeric_vals:
                    scale = detect_rating_scale(col_values)
                    result["ratings"][str(header)] = {
                        "avg": round(mean(numeric_vals), 2),
                        "min": min(numeric_vals),
                        "max": max(numeric_vals),
                        "count": len(numeric_vals),
                        "scale": scale
                    }

            # Check if it's a text feedback column
            elif is_text_column(header, col_values):
                for v in col_values:
                    if v is not None:
                        s = str(v).strip()
                        if is_meaningful_feedback(s):
                            result["text_feedback"].append({
                                "question": str(header),
                                "response": s
                            })

        # Store a sample of raw data (first 3 rows)
        for row in data_rows[:3]:
            sample = {}
            for idx, h in enumerate(headers):
                if h and idx < len(row) and row[idx] is not None:
                    sample[str(h)] = str(row[idx])
            if sample:
                result["raw_data_sample"].append(sample)

    except Exception as e:
        result["error"] = str(e)

    # Categorize
    result["category"] = categorize_session(filename)

    return result


def count_records(sessions):
    """Count files and extracted observations for one explicitly defined population."""
    return {
        "file_count": len(sessions),
        "response_rows": sum(int(session.get("response_count", 0)) for session in sessions),
        "qualitative_entries": sum(
            len(session.get("text_feedback", [])) for session in sessions
        ),
        "rating_question_aggregates": sum(
            len(session.get("ratings", {})) for session in sessions
        ),
        "rating_observations": sum(
            int(rating.get("count", 0))
            for session in sessions
            for rating in session.get("ratings", {}).values()
        ),
    }


def group_record_counts(sessions, key_function):
    """Return deterministic population counts grouped by a session attribute."""
    groups = defaultdict(list)
    for session in sessions:
        groups[str(key_function(session))].append(session)
    return {key: count_records(groups[key]) for key in sorted(groups)}


def session_year(session):
    """Return a four-digit year when available, otherwise ``unknown``."""
    date_str = str(session.get("date") or "unknown")
    if date_str != "unknown" and re.match(r"^\d{4}", date_str):
        return date_str[:4]
    return "unknown"


def rating_inventory(sessions):
    """Describe rating coverage without combining semantically different questions."""
    by_scale = defaultdict(lambda: {"question_level_aggregates": 0, "rating_observations": 0})
    for session in sessions:
        for rating in session.get("ratings", {}).values():
            scale = str(rating.get("scale") or "unknown")
            by_scale[scale]["question_level_aggregates"] += 1
            by_scale[scale]["rating_observations"] += int(rating.get("count", 0))

    counts = count_records(sessions)
    return {
        "question_level_aggregates": counts["rating_question_aggregates"],
        "rating_observations": counts["rating_observations"],
        "by_scale": {key: by_scale[key] for key in sorted(by_scale)},
        "aggregation_policy": (
            "No global average is reported. The extracted rating questions measure different "
            "constructs on 5-point and 10-point scales; normalizing and averaging them would not "
            "produce a defensible overall score."
        ),
    }


def select_top_quotes(sessions):
    """Select substantive excerpts from the post-event analysis population only."""
    all_text_feedback = []
    for session in sessions:
        for feedback in session.get("text_feedback", []):
            all_text_feedback.append(
                {
                    "session": session.get("session_name", session.get("filename", "unknown")),
                    "date": session.get("date", "unknown"),
                    "question": feedback.get("question", ""),
                    "response": feedback.get("response", ""),
                }
            )

    # Select top quotes (longer, more substantive feedback)
    # Filter out trivial responses and email-like content
    substantive_quotes = []
    for fb in all_text_feedback:
        resp = fb["response"]
        lower_resp = resp.lower().strip()
        # Skip trivial or non-feedback text
        if len(resp) < 30:
            continue
        trivial_prefixes = (
            "yes",
            "no",
            "na",
            "n/a",
            "none",
            "nothing",
            "good",
            "ok",
            "nil",
        )
        if lower_resp.startswith(trivial_prefixes):
            continue
        # Skip email-like content (contains multiple @philips.com references)
        if resp.count('@philips.com') > 1:
            continue
        # Skip content that looks like meeting notes / forwards
        if 'from:' in lower_resp and 'sent:' in lower_resp:
            continue
        substantive_quotes.append(fb)

    # Sort by a quality heuristic: prefer medium-length, diverse sessions
    # Longer is not always better (email chains are long); prefer 30-300 char range
    def quote_quality(fb):
        length = len(fb["response"])
        # Sweet spot is 50-300 characters
        if 50 <= length <= 300:
            return 3
        elif 30 <= length <= 500:
            return 2
        else:
            return 1

    substantive_quotes.sort(
        key=lambda item: (
            -quote_quality(item),
            -len(item["response"]),
            item["session"],
            item["date"],
            item["question"],
            item["response"],
        )
    )

    # Pick diverse quotes across sessions (max 3 per session)
    selected_quotes = []
    session_counts = {}
    for fb in substantive_quotes:
        session_key = fb["session"]
        if session_counts.get(session_key, 0) >= 3:
            continue
        selected_quotes.append(fb)
        session_counts[session_key] = session_counts.get(session_key, 0) + 1
        if len(selected_quotes) >= 50:
            break

    return selected_quotes


def build_before_after_comparison(sessions):
    """Describe the JSCPD pair without asserting an unsupported causal outcome."""
    before_session = None
    after_session = None
    for session in sessions:
        filename = session.get("filename", "").lower()
        if "before elimination" in filename:
            before_session = session
        elif "after elimination" in filename:
            after_session = session

    if not before_session or not after_session:
        return {}

    def comparison_record(session):
        counts = count_records([session])
        return {
            "filename": session["filename"],
            "response_rows": counts["response_rows"],
            "qualitative_entries": counts["qualitative_entries"],
            "rating_question_aggregates": counts["rating_question_aggregates"],
            "rating_observations": counts["rating_observations"],
            "ratings": session.get("ratings", {}),
        }

    return {
        "before": comparison_record(before_session),
        "after": comparison_record(after_session),
        "interpretation": (
            "The files are paired by their before/after filenames. The committed extract contains "
            "16 baseline rows and six post-exercise feedback rows, but it does not contain matched "
            "before/after outcome measures. The pair supports reporting the post-exercise ratings; "
            "it does not by itself establish code improvement or causality."
        ),
    }


def compute_summary(all_sessions):
    """Compute transparent statistics for mutually exclusive analysis populations."""
    duplicate_files = [session for session in all_sessions if session.get("is_duplicate")]
    unique_files = [session for session in all_sessions if not session.get("is_duplicate")]
    pre_session_surveys = [
        session for session in unique_files if session.get("is_pre_session_survey")
    ]
    post_event_feedback = [
        session for session in unique_files if not session.get("is_pre_session_survey")
    ]

    duplicate_counts = count_records(duplicate_files)
    duplicate_counts["filenames"] = sorted(session["filename"] for session in duplicate_files)
    pre_survey_counts = count_records(pre_session_surveys)
    pre_survey_counts["filenames"] = sorted(
        session["filename"] for session in pre_session_surveys
    )

    summary = {
        "schema_version": 2,
        "generated_from": "all_sessions_data.json",
        "population_definitions": {
            "source_files": "Every committed workbook extract, including known duplicate files.",
            "duplicates_excluded": (
                "Files marked is_duplicate=true; excluded before defining analysis populations."
            ),
            "pre_session_surveys": (
                "Unique files marked is_pre_session_survey=true; reported separately and excluded "
                "from post-event aggregates."
            ),
            "post_event_feedback": (
                "Unique files not marked as pre-session surveys. This operational population also "
                "contains interaction feedback, connect logs, and the JSCPD before/after pair, so "
                "file_count is not a count of facilitated sessions."
            ),
        },
        "source_inventory": {
            "source_files": count_records(all_sessions),
            "duplicates_excluded": duplicate_counts,
            "unique_files": count_records(unique_files),
        },
        "analysis_populations": {
            "post_event_feedback": count_records(post_event_feedback),
            "pre_session_surveys": pre_survey_counts,
        },
        "post_event_by_year": group_record_counts(post_event_feedback, session_year),
        "post_event_by_category": group_record_counts(
            post_event_feedback, lambda session: session.get("category") or "unknown"
        ),
        "rating_inventory": rating_inventory(post_event_feedback),
        "top_quotes": select_top_quotes(post_event_feedback),
        "before_after_comparison": build_before_after_comparison(post_event_feedback),
        "sessions_with_errors": sorted(
            [
                {
                    "filename": session["filename"],
                    "error": session["error"],
                }
                for session in all_sessions
                if session.get("error")
            ],
            key=lambda item: item["filename"],
        ),
        "methodology_notes": [
            "response_rows count spreadsheet rows, not verified unique participants or attendees.",
            "qualitative_entries count extracted text_feedback records; a row can contribute "
            "more than one entry, and the legacy extraction has not been manually coded or "
            "deduplicated.",
            "rating_question_aggregates count question-level summaries in the committed extract; "
            "rating_observations sum each summary's count field.",
            "The source workbooks are not committed. This summary is reproducible from the "
            "committed all_sessions_data.json extract, but participant-level rating values cannot "
            "be audited from that extract because it retains aggregates and only a three-row raw "
            "sample.",
        ],
    }

    return summary


def write_json(data, output_path):
    """Write stable, human-reviewable JSON with a trailing newline."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(data, indent=2, ensure_ascii=False, default=str) + "\n",
        encoding="utf-8",
    )


def regenerate_summary_from_json(input_path, output_path):
    """Regenerate summary statistics from the committed extracted-session JSON."""
    all_sessions = json.loads(input_path.read_text(encoding="utf-8"))
    if not isinstance(all_sessions, list):
        raise ValueError(f"Expected a JSON list of session records in {input_path}")
    summary = compute_summary(all_sessions)
    write_json(summary, output_path)
    return summary


def build_argument_parser():
    parser = argparse.ArgumentParser(
        description="Extract professional-session feedback or regenerate its summary."
    )
    parser.add_argument(
        "--from-json",
        type=Path,
        help="Regenerate only summary_stats.json from an existing all_sessions_data.json.",
    )
    parser.add_argument(
        "--summary-output",
        type=Path,
        help="Summary JSON destination (defaults to <output-dir>/summary_stats.json).",
    )
    parser.add_argument(
        "--feedback-dir",
        type=Path,
        default=FEEDBACK_DIR,
        help="Directory containing the source XLSX files.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=OUTPUT_DIR,
        help="Directory for all_sessions_data.json and summary_stats.json.",
    )
    return parser


def print_summary(summary):
    post_event = summary["analysis_populations"]["post_event_feedback"]
    pre_surveys = summary["analysis_populations"]["pre_session_surveys"]
    duplicates = summary["source_inventory"]["duplicates_excluded"]

    print("\n" + "=" * 60)
    print("PROCESSING COMPLETE")
    print("=" * 60)
    print(f"Source files: {summary['source_inventory']['source_files']['file_count']}")
    print(
        "Post-event/interaction feedback: "
        f"{post_event['file_count']} files, {post_event['response_rows']} response rows"
    )
    print(
        f"Pre-session surveys: {pre_surveys['file_count']} files, "
        f"{pre_surveys['response_rows']} response rows"
    )
    print(
        f"Duplicates excluded: {duplicates['file_count']} files, "
        f"{duplicates['response_rows']} response rows"
    )
    print(
        "Rating coverage: "
        f"{post_event['rating_question_aggregates']} question-level aggregates, "
        f"{post_event['rating_observations']} observations"
    )
    print("No heterogeneous global rating is calculated.")


def main(argv=None):
    args = build_argument_parser().parse_args(argv)
    output_dir = args.output_dir
    summary_file = args.summary_output or output_dir / "summary_stats.json"

    if args.from_json:
        summary = regenerate_summary_from_json(args.from_json, summary_file)
        print(f"Written summary stats to: {summary_file}")
        print_summary(summary)
        return

    feedback_dir = args.feedback_dir

    # Get all xlsx files
    xlsx_files = sorted(feedback_dir.glob("*.xlsx"))
    print(f"Found {len(xlsx_files)} Excel files to process")

    all_sessions = []
    for i, filepath in enumerate(xlsx_files):
        print(f"Processing [{i+1}/{len(xlsx_files)}]: {filepath.name}")
        session_data = process_file(str(filepath))
        all_sessions.append(session_data)

    # Write all sessions data
    output_file = output_dir / "all_sessions_data.json"
    write_json(all_sessions, output_file)
    print(f"\nWritten all sessions data to: {output_file}")

    # Compute and write summary
    summary = compute_summary(all_sessions)
    write_json(summary, summary_file)
    print(f"Written summary stats to: {summary_file}")
    print_summary(summary)
    if summary['sessions_with_errors']:
        print(f"Files with errors: {len(summary['sessions_with_errors'])}")
        for err in summary['sessions_with_errors']:
            print(f"  - {err['filename']}: {err['error']}")


if __name__ == "__main__":
    main()
