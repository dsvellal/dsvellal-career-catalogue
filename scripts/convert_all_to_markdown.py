#!/usr/bin/env python3
"""
Comprehensive Career Artifact to Markdown Converter
Converts PDF, DOCX, XLSX, PPTX, EML, and creates companion Markdown for images.
Preserves syntax, structure, tables, and embeds images.
"""

import os
import sys
import subprocess
import email
from email import policy
from pathlib import Path
import docx
import openpyxl
import pptx

WORKSPACE_ROOT = Path('/Users/dsvellal/Downloads/Website')

def convert_docx(file_path: Path) -> Path:
    target_md = file_path.with_suffix('.md')
    try:
        doc = docx.Document(file_path)
        lines = []
        lines.append(f"# {file_path.stem.replace('_', ' ')}\n")
        lines.append(f"> Converted from `{file_path.name}`\n")
        
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            style_name = p.style.name.lower()
            if 'heading 1' in style_name:
                lines.append(f"\n## {text}\n")
            elif 'heading 2' in style_name:
                lines.append(f"\n### {text}\n")
            elif 'heading 3' in style_name:
                lines.append(f"\n#### {text}\n")
            elif 'list' in style_name or p.text.startswith('•') or p.text.startswith('-'):
                clean_text = text.lstrip('•- ').strip()
                lines.append(f"- {clean_text}")
            else:
                lines.append(f"\n{text}\n")
                
        for t_idx, table in enumerate(doc.tables):
            lines.append(f"\n### Table {t_idx + 1}\n")
            rows = []
            for r in table.rows:
                row_cells = [c.text.strip().replace('\n', ' ') for c in r.cells]
                rows.append(row_cells)
            if rows:
                header = rows[0]
                lines.append("| " + " | ".join(header) + " |")
                lines.append("| " + " | ".join(["---"] * len(header)) + " |")
                for r in rows[1:]:
                    lines.append("| " + " | ".join(r) + " |")
                lines.append("")
                
        target_md.write_text("\n".join(lines), encoding='utf-8')
        return target_md
    except Exception as e:
        print(f"Error converting DOCX {file_path}: {e}")
        return None

def convert_xlsx(file_path: Path) -> Path:
    target_md = file_path.with_suffix('.md')
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        lines = []
        lines.append(f"# {file_path.stem.replace('_', ' ')}\n")
        lines.append(f"> Converted from spreadsheet `{file_path.name}`\n")
        
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            lines.append(f"\n## Sheet: {sheet_name}\n")
            rows_data = []
            for row in sheet.iter_rows(values_only=True):
                # Filter out completely empty rows
                if any(v is not None and str(v).strip() != '' for v in row):
                    clean_row = [str(v).replace('\n', ' ').strip() if v is not None else '' for v in row]
                    # Trim trailing empty cells
                    while clean_row and clean_row[-1] == '':
                        clean_row.pop()
                    if clean_row:
                        rows_data.append(clean_row)
            if rows_data:
                # Normalize column count to max cols
                max_cols = max(len(r) for r in rows_data)
                for r in rows_data:
                    while len(r) < max_cols:
                        r.append('')
                header = rows_data[0]
                lines.append("| " + " | ".join(header) + " |")
                lines.append("| " + " | ".join(["---"] * len(header)) + " |")
                for r in rows_data[1:250]: # cap sheet rows at 250 for readability
                    lines.append("| " + " | ".join(r) + " |")
                if len(rows_data) > 250:
                    lines.append(f"\n*...and {len(rows_data) - 250} more rows.*\n")
                lines.append("")
            else:
                lines.append("*(Empty sheet)*\n")
                
        target_md.write_text("\n".join(lines), encoding='utf-8')
        return target_md
    except Exception as e:
        print(f"Error converting XLSX {file_path}: {e}")
        return None

def convert_pdf(file_path: Path) -> Path:
    target_md = file_path.with_suffix('.md')
    try:
        res = subprocess.run(['/opt/homebrew/bin/pdftotext', str(file_path), '-'], capture_output=True, text=True)
        if res.returncode == 0:
            raw_text = res.stdout
            lines = []
            lines.append(f"# {file_path.stem.replace('_', ' ')}\n")
            lines.append(f"> Converted from document `{file_path.name}`\n")
            lines.append(raw_text)
            target_md.write_text("\n".join(lines), encoding='utf-8')
            return target_md
        else:
            print(f"pdftotext failed for {file_path}: {res.stderr}")
            return None
    except Exception as e:
        print(f"Error converting PDF {file_path}: {e}")
        return None

def convert_pptx(file_path: Path) -> Path:
    target_md = file_path.with_suffix('.md')
    try:
        prs = pptx.Presentation(file_path)
        lines = []
        lines.append(f"# {file_path.stem.replace('_', ' ')}\n")
        lines.append(f"> Converted from presentation `{file_path.name}`\n")
        
        for slide_idx, slide in enumerate(prs.slides):
            title = f"Slide {slide_idx + 1}"
            if slide.shapes.title and slide.shapes.title.text.strip():
                title = f"Slide {slide_idx + 1}: {slide.shapes.title.text.strip()}"
            lines.append(f"\n## {title}\n")
            
            for shape in slide.shapes:
                if shape.has_text_frame and shape != slide.shapes.title:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            lines.append(f"- {text}")
                            
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    lines.append(f"\n> **Notes:** {notes}\n")
            lines.append("")
            
        target_md.write_text("\n".join(lines), encoding='utf-8')
        return target_md
    except Exception as e:
        print(f"Error converting PPTX {file_path}: {e}")
        return None

def convert_eml(file_path: Path) -> Path:
    target_md = file_path.with_suffix('.md')
    try:
        with open(file_path, 'rb') as f:
            msg = email.message_from_binary_file(f, policy=policy.default)
            
        lines = []
        subject = msg.get('Subject', file_path.stem)
        from_hdr = msg.get('From', 'Unknown Sender')
        to_hdr = msg.get('To', 'Unknown Recipient')
        date_hdr = msg.get('Date', 'Unknown Date')
        
        lines.append(f"# {subject}\n")
        lines.append(f"**From:** {from_hdr}  ")
        lines.append(f"**To:** {to_hdr}  ")
        lines.append(f"**Date:** {date_hdr}  \n")
        lines.append("---\n")
        
        # Body extraction
        body = msg.get_body(preferencelist=('plain', 'html'))
        if body:
            content = body.get_content()
            lines.append(content)
        else:
            lines.append("*(No text body found)*")
            
        target_md.write_text("\n".join(lines), encoding='utf-8')
        return target_md
    except Exception as e:
        print(f"Error converting EML {file_path}: {e}")
        return None

def create_image_companion(file_path: Path) -> Path:
    target_md = file_path.with_suffix('.md')
    if target_md.exists():
        return target_md
    try:
        title = file_path.stem.replace('-', ' ').replace('_', ' ')
        lines = [
            f"# Evidence: {title}\n",
            f"> Verified Artifact Evidence: `{file_path.name}`\n",
            f"![{title}](./{file_path.name})\n",
            "## Metadata",
            f"- **Filename:** `{file_path.name}`",
            f"- **Type:** Image Verification Evidence",
            f"- **Directory:** `{file_path.parent.name}`\n"
        ]
        target_md.write_text("\n".join(lines), encoding='utf-8')
        return target_md
    except Exception as e:
        print(f"Error creating companion for image {file_path}: {e}")
        return None

def main():
    print(f"==> Scanning workspace: {WORKSPACE_ROOT}")
    
    # 5-Pillar Executive Taxonomy
    target_dirs = [
        WORKSPACE_ROOT / '01_Career_Eras',
        WORKSPACE_ROOT / '02_Evidence_and_Feedback',
        WORKSPACE_ROOT / '03_Executive_Givebacks',
        WORKSPACE_ROOT / '04_Architecture_and_Strategy',
        WORKSPACE_ROOT / '05_Credentials_and_Academics'
    ]
    
    stats = {'docx': 0, 'xlsx': 0, 'pdf': 0, 'pptx': 0, 'eml': 0, 'image': 0}
    
    for t_dir in target_dirs:
        if not t_dir.exists():
            continue
        for root, _, files in os.walk(t_dir):
            r_path = Path(root)
            # Skip node_modules and .git
            if 'node_modules' in r_path.parts or '.git' in r_path.parts or 'web' in r_path.parts or '.agents' in r_path.parts:
                continue
                
            for file_name in files:
                f_path = r_path / file_name
                ext = f_path.suffix.lower()
                
                if ext == '.docx' and not file_name.startswith('~$'):
                    if convert_docx(f_path):
                        stats['docx'] += 1
                elif ext == '.xlsx' and not file_name.startswith('~$'):
                    if convert_xlsx(f_path):
                        stats['xlsx'] += 1
                elif ext == '.pdf':
                    if convert_pdf(f_path):
                        stats['pdf'] += 1
                elif ext == '.pptx' and not file_name.startswith('~$'):
                    if convert_pptx(f_path):
                        stats['pptx'] += 1
                elif ext == '.eml':
                    if convert_eml(f_path):
                        stats['eml'] += 1
                elif ext in ['.png', '.jpg', '.jpeg']:
                    if create_image_companion(f_path):
                        stats['image'] += 1

    print("==> Conversion Complete! Summary:")
    for k, v in stats.items():
        print(f"  - {k.upper()}: {v} converted to Markdown")

if __name__ == '__main__':
    main()
